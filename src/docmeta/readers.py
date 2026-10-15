"""Turn an uploaded file into something an LLM can read.

* .docx  -> plain text (body paragraphs, tables, headers/footers)
* images -> passed to the model as vision input (no separate OCR dependency)
"""
from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from .models import ImagePart, LoadedDocument

DOCX_EXTENSIONS = {".docx"}
IMAGE_MEDIA_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".gif": "image/gif",
}
SUPPORTED_EXTENSIONS = DOCX_EXTENSIONS | set(IMAGE_MEDIA_TYPES)

MAX_IMAGE_SIDE = 3000          # px; larger scans are downscaled
MAX_IMAGE_BYTES = 4_500_000    # stay under the API's 5 MB image limit


class UnsupportedFormat(ValueError):
    pass


def load_document(path: str | Path) -> LoadedDocument:
    path = Path(path)
    return load_bytes(path.name, path.read_bytes())


def load_bytes(filename: str, data: bytes) -> LoadedDocument:
    suffix = Path(filename).suffix.lower()
    stem = Path(filename).stem
    if suffix in DOCX_EXTENSIONS:
        try:
            return LoadedDocument(name=stem, text=docx_to_text(io.BytesIO(data)))
        except Exception as exc:  # corrupt / not a real docx
            raise UnsupportedFormat(f"Could not read {filename} as .docx: {exc}") from exc
    if suffix in IMAGE_MEDIA_TYPES:
        return LoadedDocument(name=stem, images=[_prepare_image(data, IMAGE_MEDIA_TYPES[suffix])])
    raise UnsupportedFormat(
        f"Unsupported file type '{suffix}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
    )


def docx_to_text(fileobj) -> str:
    doc = Document(fileobj)
    lines: list[str] = []

    for block in doc.iter_inner_content():
        if isinstance(block, Paragraph):
            if block.text.strip():
                lines.append(block.text.strip())
        elif isinstance(block, Table):
            for row in block.rows:
                cells: list[str] = []
                for cell in row.cells:
                    text = " ".join(cell.text.split())
                    if text and (not cells or cells[-1] != text):  # merged cells repeat
                        cells.append(text)
                if cells:
                    lines.append(" | ".join(cells))

    for section in doc.sections:
        for part in (section.header, section.footer):
            for p in part.paragraphs:
                if p.text.strip():
                    lines.append(p.text.strip())

    return "\n".join(lines)


def _prepare_image(data: bytes, media_type: str) -> ImagePart:
    from PIL import Image

    try:
        img = Image.open(io.BytesIO(data))
        img.load()
    except Exception as exc:
        raise UnsupportedFormat(f"Could not read image: {exc}") from exc

    longest = max(img.size)
    if longest <= MAX_IMAGE_SIDE and len(data) <= MAX_IMAGE_BYTES:
        return ImagePart(media_type, data)

    if longest > MAX_IMAGE_SIDE:
        scale = MAX_IMAGE_SIDE / longest
        img = img.resize((int(img.width * scale), int(img.height * scale)), Image.LANCZOS)
    buf = io.BytesIO()
    img.convert("RGB").save(buf, "JPEG", quality=90)
    return ImagePart("image/jpeg", buf.getvalue())


def ocr_images(images) -> str:
    """OCR scanned pages to text, for text-only models (`pip install docmeta[ocr]`
    plus the Tesseract binary)."""
    try:
        import pytesseract
    except ImportError as exc:
        raise RuntimeError(
            "OCR needs `pip install 'docmeta[ocr]'` and the Tesseract binary "
            "(e.g. `apt install tesseract-ocr` or `brew install tesseract`)."
        ) from exc
    from PIL import Image

    pages = [pytesseract.image_to_string(Image.open(io.BytesIO(i.data))) for i in images]
    return "\n\n".join(p.strip() for p in pages if p.strip())
