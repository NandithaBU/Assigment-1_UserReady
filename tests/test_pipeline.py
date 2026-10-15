import io
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from docx import Document
from PIL import Image

from docmeta.examples import load_examples
from docmeta.pipeline import MetadataExtractor, predict_directory
from docmeta.prompts import build_user_text
from docmeta.readers import UnsupportedFormat, load_bytes, load_document


class FakeBackend:
    def __init__(self, raw=None, fail_on=None):
        self.raw = raw or {
            "agreement_value": 6500, "agreement_start_date": "20.05.2007",
            "agreement_end_date": "20.05.2008", "renewal_notice_days": None,
            "party_one": "Owner  One", "party_two": "Tenant Two",
        }
        self.calls = []
        self.fail_on = fail_on

    def extract(self, system, text, images, schema):
        self.calls.append((system, text, list(images), schema))
        if self.fail_on and self.fail_on in text:
            raise RuntimeError("boom")
        return dict(self.raw)


def make_docx(path, body="Rent is 6500", table=False):
    d = Document()
    d.add_paragraph(body)
    if table:
        t = d.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text = "Rent"
        t.rows[0].cells[1].text = "6500"
    d.save(path)


def make_png(path, size=(50, 40)):
    Image.new("RGB", size, "white").save(path)


class ReaderTests(unittest.TestCase):
    def test_docx_includes_paragraphs_and_tables(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "a.docx"
            make_docx(p, "Hello agreement", table=True)
            doc = load_document(p)
        self.assertEqual(doc.name, "a")
        self.assertIn("Hello agreement", doc.text)
        self.assertIn("Rent | 6500", doc.text)
        self.assertEqual(doc.images, [])

    def test_image_passthrough_and_downscale(self):
        with tempfile.TemporaryDirectory() as tmp:
            small, big = Path(tmp) / "s.png", Path(tmp) / "b.png"
            make_png(small)
            make_png(big, size=(4000, 3000))
            s, b = load_document(small), load_document(big)
        self.assertEqual(s.images[0].media_type, "image/png")
        self.assertEqual(b.images[0].media_type, "image/jpeg")
        img = Image.open(io.BytesIO(b.images[0].data))
        self.assertEqual(max(img.size), 3000)

    def test_unsupported_and_corrupt(self):
        with self.assertRaises(UnsupportedFormat):
            load_bytes("x.txt", b"hi")
        with self.assertRaises(UnsupportedFormat):
            load_bytes("x.docx", b"not a docx")
        with self.assertRaises(UnsupportedFormat):
            load_bytes("x.png", b"not a png")


class ExtractorTests(unittest.TestCase):
    def test_extract_normalises_and_uses_schema(self):
        backend = FakeBackend()
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "a.docx"
            make_docx(p)
            result = MetadataExtractor(backend).extract_file(p)
        self.assertEqual(result["party_one"], "Owner One")
        self.assertIsNone(result["renewal_notice_days"])
        self.assertIn("agreement_value", backend.calls[0][3]["properties"])

    def test_images_reach_backend(self):
        backend = FakeBackend()
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "a.png"
            make_png(p)
            MetadataExtractor(backend).extract_file(p)
        self.assertEqual(len(backend.calls[0][2]), 1)

    def test_prompt_contains_examples(self):
        from docmeta.models import Example, LoadedDocument
        text = build_user_text(LoadedDocument("x", text="BODY"),
                               [Example("EX TEXT", {"party_one": "Z"})])
        self.assertIn("EX TEXT", text)
        self.assertIn('"party_one": "Z"', text)
        self.assertIn("BODY", text)


class PredictTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        make_docx(self.dir / "doc1.docx", "one")
        make_docx(self.dir / "doc2.docx", "two FAILME")
        make_png(self.dir / "scan1.png")
        (self.dir / "notes.txt").write_text("ignored")

    def tearDown(self):
        self.tmp.cleanup()

    def test_predict_all_files_and_failure_isolation(self):
        ex = MetadataExtractor(FakeBackend(fail_on="FAILME"))
        out = self.dir / "out" / "pred.csv"
        df = predict_directory(ex, self.dir, out, workers=2)
        self.assertEqual(set(df["File Name"]), {"doc1", "doc2", "scan1"})
        failed = df[df["File Name"] == "doc2"].iloc[0]
        self.assertEqual(failed["Party One"], "")
        ok = df[df["File Name"] == "doc1"].iloc[0]
        self.assertEqual(ok["Aggrement Value"], "6500")
        self.assertTrue(out.exists())

    def test_manifest_restricts_and_orders(self):
        manifest = self.dir / "test.csv"
        pd.DataFrame({"File Name": ["scan1", "doc1", "missing"]}).to_csv(manifest, index=False)
        df = predict_directory(MetadataExtractor(FakeBackend()), self.dir, manifest_csv=manifest)
        self.assertEqual(list(df["File Name"]), ["scan1", "doc1", "missing"])
        self.assertEqual(df.iloc[2]["Party One"], "")

    def test_load_examples(self):
        train = self.dir / "train.csv"
        pd.DataFrame({
            "File Name": ["doc1", "doc2", "scan1"],
            "Aggrement Value": ["6500", "4000", "1"],
            "Aggrement Start Date": ["01.04.2010 ", "x", "x"],
            "Aggrement End Date": ["a", "b", "c"],
            "Renewal Notice (Days)": ["15", None, "3"],
            "Party One": ["A ", "B", "C"], "Party Two": ["D", "E", "F"],
        }).to_csv(train, index=False)
        ex = load_examples(train, self.dir, k=5)
        self.assertEqual(len(ex), 2)  # image-only rows are skipped
        self.assertEqual(ex[0].labels["agreement_value"], 6500)
        self.assertEqual(ex[0].labels["party_one"], "A")
        self.assertIsNone(ex[1].labels["renewal_notice_days"])


if __name__ == "__main__":
    unittest.main()
